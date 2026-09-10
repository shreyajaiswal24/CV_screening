"""Stage 1 - INGEST (code). File in, text out, plus a content hash.

Text is extracted locally rather than sending the file to the model, because
the raw text is needed to verify quotes against. That is a deliberate trade-off:
it costs a dependency and buys the system's central feature.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class IngestError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    try:
        reader = PdfReader(str(path))
    except Exception as e:
        raise IngestError(
            "UNREADABLE_FILE",
            f"Could not open '{path.name}' as a PDF. The file may be damaged. ({e})",
        )
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def _docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise IngestError(
            "MISSING_DEPENDENCY",
            "Word files need an extra package. Run: pip install python-docx",
        )
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:                       # CVs often hide dates in tables
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_text(path: str | Path) -> tuple[str, str, dict]:
    """Return (text, sha256_of_text, metadata). Raises IngestError with a
    message a non-developer can act on."""
    path = Path(path)
    if not path.exists():
        raise IngestError("FILE_NOT_FOUND", f"Can't find the file '{path.name}'.")

    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise IngestError(
            "UNSUPPORTED_FORMAT",
            f"'{path.name}' is a {ext or 'unknown'} file. "
            f"Supported formats are PDF, Word (.docx) and plain text.",
        )

    if ext == ".pdf":
        text = _pdf(path)
    elif ext == ".docx":
        text = _docx(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    text = text.replace("\x00", "").strip()
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    meta = {"filename": path.name, "format": ext.lstrip("."),
            "chars": len(text), "bytes": path.stat().st_size}
    return text, digest, meta
